import argparse
import getpass
import io
import os
import secrets
import sqlite3
import tempfile
from datetime import date, datetime
from functools import wraps
from pathlib import Path

from flask import Flask, g, jsonify, request, send_file, session
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from werkzeug.security import check_password_hash, generate_password_hash
from zipfile import BadZipFile


BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY") or secrets.token_hex(32),
    DATABASE=os.environ.get("DATABASE_PATH", str(BASE_DIR / "ruangsiswa.sqlite3")),
    MAX_CONTENT_LENGTH=50 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE") == "1",
    PERMANENT_SESSION_LIFETIME=60 * 60 * 8,
)
ROLES = {"admin", "guru", "siswa"}
SCHEMA_READY = False
CLASS_LEVELS = {"I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII"}
STUDENT_IMPORT_HEADERS = ("Nama Lengkap", "NIS", "Kelas", "Jenis Kelamin", "Nama Wali", "Telepon Wali")


def infer_class_level(class_name):
    first_word = class_name.strip().split(maxsplit=1)[0].upper()
    if first_word in CLASS_LEVELS or first_word.isdigit():
        return first_word
    return ""


def get_db():
    if "db" not in g:
        Path(app.config["DATABASE"]).parent.mkdir(parents=True, exist_ok=True)
        g.db = sqlite3.connect(app.config["DATABASE"])
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


def init_db():
    global SCHEMA_READY
    if SCHEMA_READY:
        return
    db = get_db()
    had_violation_types = db.execute("SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'violation_types'").fetchone() is not None
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS students (
            id TEXT PRIMARY KEY,
            nama TEXT NOT NULL,
            nis TEXT NOT NULL UNIQUE,
            kelas TEXT NOT NULL,
            gender TEXT NOT NULL CHECK (gender IN ('Laki-laki', 'Perempuan')),
            wali TEXT NOT NULL,
            telepon TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Aktif',
            foto BLOB,
            foto_mime TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL CHECK (role IN ('admin', 'guru', 'siswa')),
            student_id TEXT UNIQUE REFERENCES students(id) ON DELETE CASCADE,
            active INTEGER NOT NULL DEFAULT 1,
            must_change_password INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            CHECK ((role = 'siswa' AND student_id IS NOT NULL) OR (role != 'siswa' AND student_id IS NULL))
        );
        CREATE TABLE IF NOT EXISTS grades (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT NOT NULL REFERENCES students(id) ON DELETE CASCADE,
            mapel TEXT NOT NULL,
            semester TEXT NOT NULL CHECK (semester IN ('Ganjil', 'Genap')),
            jenis TEXT NOT NULL DEFAULT '',
            nilai INTEGER NOT NULL CHECK (nilai BETWEEN 0 AND 100),
            tanggal TEXT NOT NULL,
            created_by INTEGER REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT NOT NULL REFERENCES students(id) ON DELETE CASCADE,
            tanggal TEXT NOT NULL,
            waktu TEXT NOT NULL,
            status TEXT NOT NULL CHECK (status IN ('Hadir', 'Izin', 'Sakit', 'Alpa')),
            catatan TEXT NOT NULL DEFAULT '',
            created_by INTEGER REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS violations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT NOT NULL REFERENCES students(id) ON DELETE CASCADE,
            jenis TEXT NOT NULL,
            poin INTEGER NOT NULL CHECK (poin > 0),
            tanggal TEXT NOT NULL,
            tindak_lanjut TEXT NOT NULL CHECK (tindak_lanjut IN ('Ditindaklanjuti', 'Perlu tindak lanjut')),
            created_by INTEGER REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS violation_types (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            jenis TEXT NOT NULL COLLATE NOCASE UNIQUE,
            poin INTEGER NOT NULL CHECK (poin BETWEEN 1 AND 1000),
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT NOT NULL REFERENCES students(id) ON DELETE CASCADE,
            tagihan TEXT NOT NULL,
            jumlah INTEGER NOT NULL CHECK (jumlah > 0),
            tempo TEXT NOT NULL,
            status TEXT NOT NULL CHECK (status IN ('Lunas', 'Menunggu', 'Terlambat')),
            dibayar TEXT NOT NULL DEFAULT '',
            created_by INTEGER REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS teacher_assignments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            kelas TEXT NOT NULL,
            mapel TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE (teacher_id, kelas, mapel)
        );
        CREATE TABLE IF NOT EXISTS classes (
            nama TEXT PRIMARY KEY,
            tingkat TEXT NOT NULL DEFAULT '',
            tahun_ajaran TEXT NOT NULL DEFAULT '2025/2026',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS teacher_profiles (
            user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
            nama TEXT NOT NULL,
            nip TEXT NOT NULL DEFAULT '',
            telepon TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS school_settings (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            nama TEXT NOT NULL DEFAULT 'SMP SYARIF HIDAYATULLAH',
            npsn TEXT NOT NULL DEFAULT '',
            alamat TEXT NOT NULL DEFAULT '',
            telepon TEXT NOT NULL DEFAULT '',
            email TEXT NOT NULL DEFAULT '',
            logo BLOB,
            logo_mime TEXT NOT NULL DEFAULT '',
            auth_epoch TEXT NOT NULL DEFAULT '',
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE INDEX IF NOT EXISTS idx_grades_student ON grades(student_id);
        CREATE INDEX IF NOT EXISTS idx_attendance_student_date ON attendance(student_id, tanggal);
        CREATE TRIGGER IF NOT EXISTS trg_attendance_one_per_student_date
        BEFORE INSERT ON attendance
        WHEN EXISTS (
            SELECT 1 FROM attendance
            WHERE student_id = NEW.student_id AND tanggal = NEW.tanggal
        )
        BEGIN
            SELECT RAISE(ABORT, 'Attendance already recorded for this student and date');
        END;
        CREATE INDEX IF NOT EXISTS idx_violations_student ON violations(student_id);
        CREATE INDEX IF NOT EXISTS idx_payments_student ON payments(student_id);
        """
    )
    if not had_violation_types:
        db.execute(
            "INSERT OR IGNORE INTO violation_types (jenis, poin) "
            "SELECT MIN(trim(jenis)), MAX(poin) FROM violations "
            "WHERE trim(jenis) != '' GROUP BY trim(jenis) COLLATE NOCASE"
        )
    student_columns = {row[1] for row in db.execute("PRAGMA table_info(students)")}
    if "foto" not in student_columns:
        db.execute("ALTER TABLE students ADD COLUMN foto BLOB")
    if "foto_mime" not in student_columns:
        db.execute("ALTER TABLE students ADD COLUMN foto_mime TEXT NOT NULL DEFAULT ''")
    grade_columns = {row[1] for row in db.execute("PRAGMA table_info(grades)")}
    if "jenis" not in grade_columns:
        db.execute("ALTER TABLE grades ADD COLUMN jenis TEXT NOT NULL DEFAULT ''")
    user_columns = {row[1] for row in db.execute("PRAGMA table_info(users)")}
    if "must_change_password" not in user_columns:
        db.execute("ALTER TABLE users ADD COLUMN must_change_password INTEGER NOT NULL DEFAULT 0")
        db.execute("UPDATE users SET must_change_password = 1 WHERE role = 'siswa'")
    teacher_profile_indexes = db.execute("PRAGMA index_list(teacher_profiles)").fetchall()
    has_legacy_nip_unique = any(
        index["unique"] and index["origin"] == "u"
        and [column["name"] for column in db.execute(f"PRAGMA index_info('{index['name']}')")] == ["nip"]
        for index in teacher_profile_indexes
    )
    if has_legacy_nip_unique:
        db.execute("ALTER TABLE teacher_profiles RENAME TO teacher_profiles_legacy")
        db.execute(
            "CREATE TABLE teacher_profiles ("
            "user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE, "
            "nama TEXT NOT NULL, nip TEXT NOT NULL DEFAULT '', "
            "telepon TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)"
        )
        db.execute(
            "INSERT INTO teacher_profiles (user_id, nama, nip, telepon, created_at) "
            "SELECT user_id, nama, nip, telepon, created_at FROM teacher_profiles_legacy"
        )
        db.execute("DROP TABLE teacher_profiles_legacy")
    db.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_teacher_profiles_nip "
        "ON teacher_profiles(nip) WHERE nip != ''"
    )
    db.execute(
        "INSERT OR IGNORE INTO classes (nama, tingkat) "
        "SELECT DISTINCT kelas, CASE WHEN kelas GLOB '[0-9]*' THEN substr(kelas, 1, 1) ELSE '' END FROM students WHERE trim(kelas) != ''"
    )
    db.execute(
        "INSERT OR IGNORE INTO classes (nama, tingkat) "
        "SELECT DISTINCT kelas, CASE WHEN kelas GLOB '[0-9]*' THEN substr(kelas, 1, 1) ELSE '' END FROM teacher_assignments WHERE trim(kelas) != ''"
    )
    for row in db.execute("SELECT nama FROM classes WHERE tingkat = ''").fetchall():
        level = infer_class_level(row["nama"])
        if level:
            db.execute("UPDATE classes SET tingkat = ? WHERE nama = ?", (level, row["nama"]))
    db.execute("INSERT OR IGNORE INTO school_settings (id, nama) VALUES (1, 'SMA Nusantara')")
    db.execute("UPDATE school_settings SET auth_epoch = ? WHERE id = 1 AND auth_epoch = ''", (secrets.token_urlsafe(32),))
    db.commit()
    SCHEMA_READY = True


@app.before_request
def prepare_database():
    init_db()


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def current_user():
    if "current_user" not in g:
        user_id = session.get("user_id")
        g.current_user = get_db().execute(
            "SELECT id, username, role, student_id, active, must_change_password, password_hash FROM users WHERE id = ?",
            (user_id,),
        ).fetchone() if user_id else None
    return g.current_user


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user or not user["active"]:
            if session.get("user_id"):
                session.clear()
            return jsonify(error="Silakan masuk terlebih dahulu."), 401
        auth_epoch = get_db().execute("SELECT auth_epoch FROM school_settings WHERE id = 1").fetchone()[0]
        if not secrets.compare_digest(session.get("auth_epoch", ""), auth_epoch):
            session.clear()
            return jsonify(error="Sesi telah berakhir. Silakan masuk kembali."), 401
        if user["must_change_password"] and request.endpoint not in {"me", "logout", "change_password"}:
            return jsonify(error="Ganti kata sandi terlebih dahulu untuk melanjutkan.", must_change_password=True), 428
        return view(*args, **kwargs)
    return wrapped


def roles_required(*roles):
    def decorate(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            user = current_user()
            if not user or user["role"] not in roles:
                return jsonify(error="Anda tidak memiliki akses ke data ini."), 403
            return view(*args, **kwargs)
        return wrapped
    return decorate


def csrf_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        expected = session.get("csrf_token")
        supplied = request.headers.get("X-CSRF-Token", "")
        if not expected or not secrets.compare_digest(expected, supplied):
            return jsonify(error="Token keamanan tidak valid. Muat ulang halaman."), 400
        return view(*args, **kwargs)
    return wrapped


def payload():
    return request.get_json(silent=True) or {}


def text(data, key, label, maximum=160):
    value = str(data.get(key, "")).strip()
    if not value:
        raise ValueError(f"{label} wajib diisi.")
    if len(value) > maximum:
        raise ValueError(f"{label} terlalu panjang.")
    return value


def read_raster_upload(uploaded, label, maximum=2 * 1024 * 1024):
    if not uploaded or not uploaded.filename:
        return None, ""
    contents = uploaded.read(maximum + 1)
    if len(contents) > maximum:
        raise ValueError(f"Ukuran {label} maksimal 2 MB.")
    if contents.startswith(b"\x89PNG\r\n\x1a\n"):
        mime = "image/png"
    elif contents.startswith(b"\xff\xd8\xff"):
        mime = "image/jpeg"
    elif len(contents) >= 12 and contents[:4] == b"RIFF" and contents[8:12] == b"WEBP":
        mime = "image/webp"
    else:
        raise ValueError(f"Format {label} harus PNG, JPEG, atau WebP.")
    return contents, mime


def student_json(row):
    item = dict(row)
    has_photo = bool(item.pop("foto", None))
    item.pop("foto_mime", None)
    item["foto_url"] = f"/api/students/{item['id']}/photo" if has_photo else ""
    return item


def valid_date(value, label):
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label} harus menggunakan tanggal yang valid.") from None


def scoped_rows(table, order="id DESC"):
    user = current_user()
    if user["role"] == "siswa":
        if not user["student_id"]:
            return []
        return get_db().execute(
            f"SELECT * FROM {table} WHERE student_id = ? ORDER BY {order}",
            (user["student_id"],),
        ).fetchall()
    if user["role"] == "guru":
        order_clause = ", ".join(
            f"r.{part.strip().split()[0]} {part.strip().split()[-1]}" for part in order.split(",")
        )
        if table == "grades":
            return get_db().execute(
                "SELECT r.* FROM grades r JOIN students s ON s.id = r.student_id "
                "WHERE EXISTS (SELECT 1 FROM teacher_assignments a WHERE a.teacher_id = ? AND a.kelas = s.kelas AND a.mapel = r.mapel) "
                f"ORDER BY {order_clause}",
                (user["id"],),
            ).fetchall()
        if table in {"attendance", "violations"}:
            return get_db().execute(
                f"SELECT r.* FROM {table} r JOIN students s ON s.id = r.student_id "
                "WHERE EXISTS (SELECT 1 FROM teacher_assignments a WHERE a.teacher_id = ? AND a.kelas = s.kelas) "
                f"ORDER BY {order_clause}",
                (user["id"],),
            ).fetchall()
    return get_db().execute(f"SELECT * FROM {table} ORDER BY {order}").fetchall()


def teacher_can_access_student(student_id, subject=None):
    user = current_user()
    if user["role"] != "guru":
        return True
    student_row = get_db().execute("SELECT kelas FROM students WHERE id = ?", (student_id,)).fetchone()
    if not student_row:
        return False
    if subject:
        return get_db().execute(
            "SELECT 1 FROM teacher_assignments WHERE teacher_id = ? AND kelas = ? AND mapel = ?",
            (user["id"], student_row["kelas"], subject),
        ).fetchone() is not None
    return get_db().execute(
        "SELECT 1 FROM teacher_assignments WHERE teacher_id = ? AND kelas = ? LIMIT 1",
        (user["id"], student_row["kelas"]),
    ).fetchone() is not None


def rows_as_json(rows, rename_student=False):
    result = []
    for row in rows:
        item = dict(row)
        if rename_student and "student_id" in item:
            item["siswa"] = item.pop("student_id")
        if "tindak_lanjut" in item:
            item["tindakLanjut"] = item.pop("tindak_lanjut")
        result.append(item)
    return result


@app.get("/")
def index():
    return send_file(BASE_DIR / "index.html")


@app.get("/styles.css")
def styles():
    return send_file(BASE_DIR / "styles.css")


@app.get("/frontend.js")
def frontend_script():
    return send_file(BASE_DIR / "frontend.js")


@app.get("/api/csrf")
def get_csrf():
    token = session.setdefault("csrf_token", secrets.token_urlsafe(32))
    return jsonify(csrf=token)


@app.post("/api/login")
@csrf_required
def login():
    data = payload()
    username = str(data.get("username", "")).strip()
    password = str(data.get("password", ""))
    user = get_db().execute("SELECT * FROM users WHERE username = ? AND active = 1", (username,)).fetchone()
    if not user or not check_password_hash(user["password_hash"], password):
        return jsonify(error="Nama pengguna atau kata sandi salah."), 401
    session.clear()
    session.permanent = True
    session["user_id"] = user["id"]
    session["auth_epoch"] = get_db().execute("SELECT auth_epoch FROM school_settings WHERE id = 1").fetchone()[0]
    session["csrf_token"] = secrets.token_urlsafe(32)
    return jsonify(user={"id": user["id"], "username": user["username"], "role": user["role"], "student_id": user["student_id"], "must_change_password": bool(user["must_change_password"])}, csrf=session["csrf_token"])


@app.put("/api/account/password")
@login_required
@roles_required("siswa")
@csrf_required
def change_password():
    data = payload()
    password = str(data.get("password", ""))
    confirmation = str(data.get("confirmation", ""))
    if len(password) < 8:
        return jsonify(error="Kata sandi baru minimal 8 karakter."), 400
    if password != confirmation:
        return jsonify(error="Konfirmasi kata sandi tidak cocok."), 400
    user = current_user()
    if check_password_hash(user["password_hash"], password):
        return jsonify(error="Kata sandi baru harus berbeda dari kata sandi saat ini."), 400
    db = get_db()
    db.execute("UPDATE users SET password_hash = ?, must_change_password = 0 WHERE id = ?", (generate_password_hash(password), user["id"]))
    db.commit()
    session["csrf_token"] = secrets.token_urlsafe(32)
    return jsonify(
        ok=True,
        user={"id": user["id"], "username": user["username"], "role": user["role"], "student_id": user["student_id"], "must_change_password": False},
        csrf=session["csrf_token"],
    )


@app.post("/api/logout")
@login_required
@csrf_required
def logout():
    session.clear()
    return jsonify(ok=True)


@app.get("/api/me")
@login_required
def me():
    user = current_user()
    return jsonify(user={"id": user["id"], "username": user["username"], "role": user["role"], "student_id": user["student_id"], "must_change_password": bool(user["must_change_password"])}, csrf=session["csrf_token"])


@app.get("/api/students")
@login_required
@roles_required("admin", "guru", "siswa")
def list_students():
    user = current_user()
    if user["role"] == "siswa":
        rows = get_db().execute(
            "SELECT s.*, u.username AS akun FROM students s LEFT JOIN users u ON u.student_id = s.id WHERE s.id = ?",
            (user["student_id"],),
        ).fetchall()
    else:
        rows = get_db().execute(
            "SELECT s.*, u.username AS akun FROM students s LEFT JOIN users u ON u.student_id = s.id ORDER BY s.kelas, s.nama"
        ).fetchall()
    return jsonify(items=[student_json(row) for row in rows])


@app.get("/api/students/import-template")
@login_required
@roles_required("admin")
def download_student_import_template():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data Siswa"
    worksheet.append(STUDENT_IMPORT_HEADERS)
    worksheet.freeze_panes = "A2"
    for column, width in zip("ABCDEF", (28, 18, 22, 20, 28, 22)):
        worksheet.column_dimensions[column].width = width
    instructions = workbook.create_sheet("Petunjuk")
    instructions.append(("Petunjuk impor siswa",))
    instructions.append(("Isi satu siswa per baris pada sheet Data Siswa.",))
    instructions.append(("Jenis Kelamin harus Laki-laki atau Perempuan.",))
    instructions.append(("Kelas harus sudah dibuat pada menu Data Kelas.",))
    instructions.append(("NIS harus unik; gunakan format teks untuk mempertahankan angka nol di depan.",))
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="template-impor-siswa.xlsx",
    )


@app.post("/api/students/import")
@login_required
@roles_required("admin")
@csrf_required
def import_students():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify(error="Pilih file Excel untuk diimpor."), 400
    if Path(uploaded.filename).suffix.lower() != ".xlsx":
        return jsonify(error="Format file harus .xlsx."), 400
    contents = uploaded.read(10 * 1024 * 1024 + 1)
    if len(contents) > 10 * 1024 * 1024:
        return jsonify(error="Ukuran file Excel maksimal 10 MB."), 400
    try:
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError):
        return jsonify(error="File Excel tidak dapat dibaca."), 400

    worksheet = workbook["Data Siswa"] if "Data Siswa" in workbook.sheetnames else workbook.active
    header = tuple(str(value).strip() if value is not None else "" for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    if header != STUDENT_IMPORT_HEADERS:
        workbook.close()
        return jsonify(error="Kolom Excel tidak sesuai. Unduh dan gunakan template impor siswa."), 400

    db = get_db()
    rows = []
    seen_nis = set()
    try:
        for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            if len(rows) >= 5000:
                raise ValueError("Maksimal 5.000 siswa dalam satu impor.")
            data = dict(zip(("nama", "nis", "kelas", "gender", "wali", "telepon"), cells))
            values = {
                "nama": text(data, "nama", "Nama lengkap"),
                "nis": text(data, "nis", "NIS", 40),
                "kelas": text(data, "kelas", "Kelas", 80),
                "gender": text(data, "gender", "Jenis kelamin"),
                "wali": text(data, "wali", "Nama wali"),
                "telepon": text(data, "telepon", "Nomor telepon", 40),
            }
            if values["gender"] not in {"Laki-laki", "Perempuan"}:
                raise ValueError("Jenis kelamin harus Laki-laki atau Perempuan.")
            if not db.execute("SELECT 1 FROM classes WHERE nama = ?", (values["kelas"],)).fetchone():
                raise ValueError("Kelas harus sudah terdaftar pada menu Data Kelas.")
            if values["nis"] in seen_nis or db.execute("SELECT 1 FROM students WHERE nis = ?", (values["nis"],)).fetchone():
                raise sqlite3.IntegrityError("NIS sudah digunakan.")
            seen_nis.add(values["nis"])
            rows.append((row_number, values))
    except ValueError as error:
        workbook.close()
        return jsonify(error=f"Baris {row_number}: {error}"), 400
    except sqlite3.IntegrityError:
        workbook.close()
        return jsonify(error=f"Baris {row_number}: NIS sudah digunakan."), 409
    workbook.close()
    if not rows:
        return jsonify(error="Tidak ada data siswa untuk diimpor."), 400

    try:
        for _row_number, values in rows:
            student_id = f"S-{secrets.token_hex(6)}"
            db.execute(
                "INSERT INTO students (id, nama, nis, kelas, gender, wali, telepon) "
                "VALUES (:id, :nama, :nis, :kelas, :gender, :wali, :telepon)",
                {"id": student_id, **values},
            )
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        return jsonify(error="Impor dibatalkan karena ada NIS yang sudah digunakan."), 409
    return jsonify(ok=True, imported=len(rows)), 201


@app.get("/api/students/<student_id>/photo")
@login_required
@roles_required("admin", "guru", "siswa")
def get_student_photo(student_id):
    user = current_user()
    if user["role"] == "siswa" and user["student_id"] != student_id:
        return jsonify(error="Data siswa tidak ditemukan."), 404
    row = get_db().execute("SELECT foto, foto_mime FROM students WHERE id = ?", (student_id,)).fetchone()
    if not row or not row["foto"]:
        return jsonify(error="Foto siswa belum tersedia."), 404
    response = send_file(io.BytesIO(row["foto"]), mimetype=row["foto_mime"], max_age=0)
    response.headers["Cache-Control"] = "private, no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


@app.post("/api/students")
@login_required
@roles_required("admin")
@csrf_required
def create_student():
    data = request.form if request.mimetype == "multipart/form-data" else payload()
    try:
        values = {
            "nama": text(data, "nama", "Nama lengkap"),
            "nis": text(data, "nis", "NIS", 40),
            "kelas": text(data, "kelas", "Kelas", 80),
            "gender": text(data, "gender", "Jenis kelamin"),
            "wali": text(data, "wali", "Nama wali"),
            "telepon": text(data, "telepon", "Nomor telepon", 40),
        }
        if values["gender"] not in {"Laki-laki", "Perempuan"}:
            raise ValueError("Jenis kelamin tidak valid.")
        if not get_db().execute("SELECT 1 FROM classes WHERE nama = ?", (values["kelas"],)).fetchone():
            raise ValueError("Pilih kelas yang sudah terdaftar.")
        photo, photo_mime = read_raster_upload(request.files.get("foto"), "foto siswa")
        student_id = f"S-{secrets.token_hex(6)}"
        get_db().execute(
            "INSERT INTO students (id, nama, nis, kelas, gender, wali, telepon, foto, foto_mime) VALUES (:id, :nama, :nis, :kelas, :gender, :wali, :telepon, :foto, :foto_mime)",
            {"id": student_id, **values, "foto": photo, "foto_mime": photo_mime},
        )
        get_db().commit()
        return jsonify(item={"id": student_id, **values, "status": "Aktif", "akun": None, "foto_url": f"/api/students/{student_id}/photo" if photo else ""}), 201
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="NIS sudah digunakan."), 409


@app.put("/api/students/<student_id>")
@login_required
@roles_required("admin")
@csrf_required
def update_student(student_id):
    data = request.form if request.mimetype == "multipart/form-data" else payload()
    try:
        values = {
            "nama": text(data, "nama", "Nama lengkap"),
            "nis": text(data, "nis", "NIS", 40),
            "kelas": text(data, "kelas", "Kelas", 80),
            "gender": text(data, "gender", "Jenis kelamin"),
            "wali": text(data, "wali", "Nama wali"),
            "telepon": text(data, "telepon", "Nomor telepon", 40),
        }
        if values["gender"] not in {"Laki-laki", "Perempuan"}:
            raise ValueError("Jenis kelamin tidak valid.")
        if not get_db().execute("SELECT 1 FROM classes WHERE nama = ?", (values["kelas"],)).fetchone():
            raise ValueError("Pilih kelas yang sudah terdaftar.")
        existing = get_db().execute("SELECT 1 FROM students WHERE id = ?", (student_id,)).fetchone()
        if not existing:
            return jsonify(error="Data siswa tidak ditemukan."), 404
        photo, photo_mime = read_raster_upload(request.files.get("foto"), "foto siswa")
        remove_photo = data.get("hapus_foto") == "1"
        if photo:
            get_db().execute(
                "UPDATE students SET nama = :nama, nis = :nis, kelas = :kelas, gender = :gender, wali = :wali, telepon = :telepon, foto = :foto, foto_mime = :foto_mime WHERE id = :id",
                {**values, "id": student_id, "foto": photo, "foto_mime": photo_mime},
            )
        elif remove_photo:
            get_db().execute(
                "UPDATE students SET nama = :nama, nis = :nis, kelas = :kelas, gender = :gender, wali = :wali, telepon = :telepon, foto = NULL, foto_mime = '' WHERE id = :id",
                {**values, "id": student_id},
            )
        else:
            get_db().execute(
                "UPDATE students SET nama = :nama, nis = :nis, kelas = :kelas, gender = :gender, wali = :wali, telepon = :telepon WHERE id = :id",
                {**values, "id": student_id},
            )
        get_db().commit()
        row = get_db().execute("SELECT s.*, u.username AS akun FROM students s LEFT JOIN users u ON u.student_id = s.id WHERE s.id = ?", (student_id,)).fetchone()
        return jsonify(item=student_json(row))
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="NIS sudah digunakan."), 409


@app.delete("/api/students/<student_id>")
@login_required
@roles_required("admin")
@csrf_required
def delete_student(student_id):
    cursor = get_db().execute("DELETE FROM students WHERE id = ?", (student_id,))
    get_db().commit()
    if not cursor.rowcount:
        return jsonify(error="Data siswa tidak ditemukan."), 404
    return jsonify(ok=True)


@app.get("/api/users")
@login_required
@roles_required("admin")
def list_users():
    rows = get_db().execute(
        "SELECT u.id, u.username, u.role, u.student_id, u.active, s.nama AS nama_siswa "
        "FROM users u LEFT JOIN students s ON s.id = u.student_id ORDER BY u.id"
    ).fetchall()
    return jsonify(items=[dict(row) for row in rows])


@app.post("/api/users")
@login_required
@roles_required("admin")
@csrf_required
def create_user():
    data = payload()
    try:
        role = text(data, "role", "Peran")
        username_value = str(data.get("username", "")).strip()
        password = str(data.get("password", ""))
        student_id = data.get("student_id") or None
        if role == "siswa" and not username_value:
            if not student_id:
                raise ValueError("Akun siswa harus terhubung dengan data siswa.")
            student = get_db().execute("SELECT nis FROM students WHERE id = ?", (student_id,)).fetchone()
            if not student:
                raise ValueError("Pilih data siswa yang valid.")
            username_value = student["nis"]
        username = text({"username": username_value}, "username", "Nama pengguna", 50)
        if not username.replace("_", "").replace("-", "").isalnum():
            raise ValueError("Nama pengguna hanya boleh berisi huruf, angka, tanda hubung, dan garis bawah.")
        if len(password) < 8:
            raise ValueError("Kata sandi minimal 8 karakter.")
        if role not in ROLES:
            raise ValueError("Peran akun tidak valid.")
        if role == "siswa" and not student_id:
            raise ValueError("Akun siswa harus terhubung dengan data siswa.")
        if role != "siswa" and student_id:
            raise ValueError("Hanya akun siswa yang dihubungkan ke data siswa.")
        get_db().execute(
            "INSERT INTO users (username, password_hash, role, student_id, must_change_password) VALUES (?, ?, ?, ?, ?)",
            (username, generate_password_hash(password), role, student_id, int(role == "siswa")),
        )
        get_db().commit()
        return jsonify(ok=True), 201
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Nama pengguna sudah digunakan atau data siswa sudah memiliki akun."), 409


@app.put("/api/users/<int:user_id>")
@login_required
@roles_required("admin")
@csrf_required
def update_user(user_id):
    data = payload()
    try:
        username = text(data, "username", "Nama pengguna", 50)
        role = text(data, "role", "Peran")
        if not username.replace("_", "").replace("-", "").isalnum():
            raise ValueError("Nama pengguna hanya boleh berisi huruf, angka, tanda hubung, dan garis bawah.")
        if role not in ROLES:
            raise ValueError("Peran akun tidak valid.")
        password = str(data.get("password", ""))
        if password and len(password) < 8:
            raise ValueError("Kata sandi minimal 8 karakter.")
        student_id = data.get("student_id") or None
        if role == "siswa" and not student_id:
            raise ValueError("Akun siswa harus terhubung dengan data siswa.")
        if role != "siswa" and student_id:
            raise ValueError("Hanya akun siswa yang dihubungkan ke data siswa.")

        db = get_db()
        existing = db.execute("SELECT role, active, must_change_password FROM users WHERE id = ?", (user_id,)).fetchone()
        if not existing:
            return jsonify(error="Akun pengguna tidak ditemukan."), 404
        if role == "siswa" and not db.execute("SELECT 1 FROM students WHERE id = ?", (student_id,)).fetchone():
            raise ValueError("Pilih data siswa yang valid.")
        if existing["role"] == "admin" and existing["active"] and role != "admin":
            active_admins = db.execute("SELECT COUNT(*) FROM users WHERE role = 'admin' AND active = 1").fetchone()[0]
            if active_admins <= 1:
                return jsonify(error="Admin aktif terakhir tidak dapat diubah perannya."), 409

        must_change_password = int(role == "siswa" and (existing["role"] != "siswa" or existing["must_change_password"] or bool(password)))
        if password:
            db.execute(
                "UPDATE users SET username = ?, role = ?, student_id = ?, password_hash = ?, must_change_password = ? WHERE id = ?",
                (username, role, student_id, generate_password_hash(password), must_change_password, user_id),
            )
        else:
            db.execute(
                "UPDATE users SET username = ?, role = ?, student_id = ?, must_change_password = ? WHERE id = ?",
                (username, role, student_id, must_change_password, user_id),
            )
        db.commit()
        return jsonify(ok=True)
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Nama pengguna sudah digunakan atau siswa sudah terhubung ke akun lain."), 409


@app.post("/api/users/<int:user_id>/reset-password")
@login_required
@roles_required("admin")
@csrf_required
def reset_student_password(user_id):
    db = get_db()
    account = db.execute(
        "SELECT u.role, u.student_id, s.nis FROM users u "
        "LEFT JOIN students s ON s.id = u.student_id WHERE u.id = ?",
        (user_id,),
    ).fetchone()
    if not account:
        return jsonify(error="Akun pengguna tidak ditemukan."), 404
    if account["role"] != "siswa" or not account["student_id"] or not account["nis"]:
        return jsonify(error="Reset password hanya tersedia untuk akun siswa."), 400
    db.execute(
        "UPDATE users SET password_hash = ?, must_change_password = 1 WHERE id = ?",
        (generate_password_hash(account["nis"]), user_id),
    )
    db.commit()
    return jsonify(ok=True)


@app.get("/api/classes")
@login_required
@roles_required("admin", "guru", "siswa")
def list_classes():
    rows = get_db().execute(
        "SELECT c.nama, c.tingkat, c.tahun_ajaran, "
        "(SELECT COUNT(*) FROM students s WHERE s.kelas = c.nama) AS jumlah_siswa, "
        "(SELECT COUNT(*) FROM teacher_assignments a WHERE a.kelas = c.nama) AS jumlah_pengampu "
        "FROM classes c ORDER BY c.tingkat, c.nama"
    ).fetchall()
    return jsonify(items=[dict(row) for row in rows])


@app.post("/api/classes")
@login_required
@roles_required("admin")
@csrf_required
def create_class():
    data = payload()
    try:
        name = text(data, "nama", "Nama kelas", 80)
        grade_level = str(data.get("tingkat", "")).strip()
        school_year = str(data.get("tahun_ajaran", "2025/2026")).strip()
        if len(grade_level) > 20 or len(school_year) > 20:
            raise ValueError("Tingkat atau tahun ajaran terlalu panjang.")
        get_db().execute("INSERT INTO classes (nama, tingkat, tahun_ajaran) VALUES (?, ?, ?)", (name, grade_level, school_year))
        get_db().commit()
        return jsonify(item={"nama": name, "tingkat": grade_level, "tahun_ajaran": school_year, "jumlah_siswa": 0, "jumlah_pengampu": 0}), 201
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Nama kelas sudah terdaftar."), 409


@app.put("/api/classes/<path:class_name>")
@login_required
@roles_required("admin")
@csrf_required
def update_class(class_name):
    data = payload()
    try:
        new_name = text(data, "nama", "Nama kelas", 80)
        grade_level = str(data.get("tingkat", "")).strip()
        school_year = str(data.get("tahun_ajaran", "2025/2026")).strip()
        if len(grade_level) > 20 or len(school_year) > 20:
            raise ValueError("Tingkat atau tahun ajaran terlalu panjang.")
        db = get_db()
        if not db.execute("SELECT 1 FROM classes WHERE nama = ?", (class_name,)).fetchone():
            return jsonify(error="Kelas tidak ditemukan."), 404
        if new_name != class_name and db.execute("SELECT 1 FROM classes WHERE nama = ?", (new_name,)).fetchone():
            return jsonify(error="Nama kelas sudah terdaftar."), 409
        db.execute("UPDATE students SET kelas = ? WHERE kelas = ?", (new_name, class_name))
        db.execute("UPDATE teacher_assignments SET kelas = ? WHERE kelas = ?", (new_name, class_name))
        db.execute("UPDATE classes SET nama = ?, tingkat = ?, tahun_ajaran = ? WHERE nama = ?", (new_name, grade_level, school_year, class_name))
        db.commit()
        return jsonify(ok=True)
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Perubahan nama kelas bertabrakan dengan data yang sudah ada."), 409


@app.delete("/api/classes/<path:class_name>")
@login_required
@roles_required("admin")
@csrf_required
def delete_class(class_name):
    db = get_db()
    if db.execute("SELECT 1 FROM students WHERE kelas = ? LIMIT 1", (class_name,)).fetchone():
        return jsonify(error="Kelas masih memiliki siswa; pindahkan siswa terlebih dahulu."), 409
    if db.execute("SELECT 1 FROM teacher_assignments WHERE kelas = ? LIMIT 1", (class_name,)).fetchone():
        return jsonify(error="Kelas masih memiliki penugasan guru; hapus penugasan terlebih dahulu."), 409
    cursor = db.execute("DELETE FROM classes WHERE nama = ?", (class_name,))
    db.commit()
    if not cursor.rowcount:
        return jsonify(error="Kelas tidak ditemukan."), 404
    return jsonify(ok=True)


@app.get("/api/teachers")
@login_required
@roles_required("admin")
def list_teachers():
    rows = get_db().execute(
        "SELECT u.id AS user_id, u.username, u.active, p.nama, p.nip, p.telepon, "
        "p.user_id IS NULL AS profil_belum_lengkap, "
        "(SELECT COUNT(*) FROM teacher_assignments a WHERE a.teacher_id = u.id) AS jumlah_pengampu "
        "FROM users u LEFT JOIN teacher_profiles p ON p.user_id = u.id "
        "WHERE u.role = 'guru' ORDER BY COALESCE(p.nama, u.username)"
    ).fetchall()
    return jsonify(items=[dict(row) for row in rows])


@app.put("/api/teachers/<int:user_id>")
@login_required
@roles_required("admin")
@csrf_required
def save_teacher_profile(user_id):
    data = payload()
    try:
        name = text(data, "nama", "Nama lengkap", 120)
        nip = str(data.get("nip", "")).strip()
        if len(nip) > 40:
            raise ValueError("NIP terlalu panjang.")
        phone = str(data.get("telepon", "")).strip()
        if len(phone) > 40:
            raise ValueError("Nomor telepon terlalu panjang.")
        db = get_db()
        if not db.execute("SELECT 1 FROM users WHERE id = ? AND role = 'guru'", (user_id,)).fetchone():
            return jsonify(error="Akun guru tidak ditemukan."), 404
        db.execute(
            "INSERT INTO teacher_profiles (user_id, nama, nip, telepon) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(user_id) DO UPDATE SET nama = excluded.nama, nip = excluded.nip, telepon = excluded.telepon",
            (user_id, name, nip, phone),
        )
        db.commit()
        return jsonify(ok=True)
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="NIP sudah digunakan guru lain."), 409


@app.delete("/api/teachers/<int:user_id>")
@login_required
@roles_required("admin")
@csrf_required
def delete_teacher(user_id):
    db = get_db()
    teacher = db.execute("SELECT 1 FROM users WHERE id = ? AND role = 'guru'", (user_id,)).fetchone()
    if not teacher:
        return jsonify(error="Akun guru tidak ditemukan."), 404
    db.execute("DELETE FROM users WHERE id = ?", (user_id,))
    db.commit()
    return jsonify(ok=True)


@app.get("/api/assignments")
@login_required
@roles_required("admin", "guru")
def list_teacher_assignments():
    user = current_user()
    sql = (
        "SELECT a.id, a.teacher_id, COALESCE(p.nama, u.username) AS guru, a.kelas, a.mapel "
        "FROM teacher_assignments a JOIN users u ON u.id = a.teacher_id "
        "LEFT JOIN teacher_profiles p ON p.user_id = u.id"
    )
    if user["role"] == "guru":
        rows = get_db().execute(sql + " WHERE a.teacher_id = ? ORDER BY a.kelas, a.mapel", (user["id"],)).fetchall()
    else:
        rows = get_db().execute(sql + " ORDER BY u.username, a.kelas, a.mapel").fetchall()
    return jsonify(items=[dict(row) for row in rows])


@app.post("/api/assignments")
@login_required
@roles_required("admin")
@csrf_required
def create_teacher_assignment():
    data = payload()
    try:
        teacher_id = int(data.get("teacher_id"))
    except (TypeError, ValueError):
        return jsonify(error="Pilih akun guru yang valid."), 400
    kelas = text(data, "kelas", "Kelas", 80)
    mapel = text(data, "mapel", "Mata pelajaran", 120)
    if not get_db().execute("SELECT 1 FROM classes WHERE nama = ? LIMIT 1", (kelas,)).fetchone():
        return jsonify(error="Kelas belum terdaftar pada data kelas."), 400
    teacher = get_db().execute("SELECT id FROM users WHERE id = ? AND role = 'guru' AND active = 1", (teacher_id,)).fetchone()
    if not teacher:
        return jsonify(error="Akun guru tidak ditemukan atau tidak aktif."), 400
    try:
        cursor = get_db().execute(
            "INSERT INTO teacher_assignments (teacher_id, kelas, mapel) VALUES (?, ?, ?)",
            (teacher_id, kelas, mapel),
        )
        get_db().commit()
        row = get_db().execute(
            "SELECT a.id, a.teacher_id, COALESCE(p.nama, u.username) AS guru, a.kelas, a.mapel FROM teacher_assignments a JOIN users u ON u.id = a.teacher_id LEFT JOIN teacher_profiles p ON p.user_id = u.id WHERE a.id = ?",
            (cursor.lastrowid,),
        ).fetchone()
        return jsonify(item=dict(row)), 201
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Penugasan guru tersebut sudah ada."), 409


@app.delete("/api/assignments/<int:assignment_id>")
@login_required
@roles_required("admin")
@csrf_required
def delete_teacher_assignment(assignment_id):
    cursor = get_db().execute("DELETE FROM teacher_assignments WHERE id = ?", (assignment_id,))
    get_db().commit()
    if not cursor.rowcount:
        return jsonify(error="Penugasan tidak ditemukan."), 404
    return jsonify(ok=True)


def settings_json():
    row = get_db().execute("SELECT nama, npsn, alamat, telepon, email, logo FROM school_settings WHERE id = 1").fetchone()
    return {
        "nama": row["nama"], "npsn": row["npsn"], "alamat": row["alamat"],
        "telepon": row["telepon"], "email": row["email"],
        "logo_url": "/api/settings/logo" if row["logo"] else "",
    }


@app.get("/api/public-settings")
def public_settings():
    row = get_db().execute("SELECT nama, logo FROM school_settings WHERE id = 1").fetchone()
    return jsonify(nama=row["nama"], logo_url="/api/settings/logo" if row["logo"] else "")


@app.get("/api/settings")
@login_required
@roles_required("admin")
def get_school_settings():
    return jsonify(settings_json())


@app.get("/api/settings/logo")
def get_school_logo():
    row = get_db().execute("SELECT logo, logo_mime FROM school_settings WHERE id = 1").fetchone()
    if not row["logo"]:
        return jsonify(error="Logo sekolah belum diunggah."), 404
    response = send_file(io.BytesIO(row["logo"]), mimetype=row["logo_mime"], max_age=0)
    response.headers["Cache-Control"] = "no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


@app.put("/api/settings")
@login_required
@roles_required("admin")
@csrf_required
def update_school_settings():
    data = request.form if request.mimetype == "multipart/form-data" else payload()
    try:
        name = text(data, "nama", "Nama sekolah", 120)
        npsn = str(data.get("npsn", "")).strip()
        address = str(data.get("alamat", "")).strip()
        phone = str(data.get("telepon", "")).strip()
        email = str(data.get("email", "")).strip()
        if len(npsn) > 20 or (npsn and not npsn.isdigit()):
            raise ValueError("NPSN harus berupa angka.")
        if len(address) > 500 or len(phone) > 40 or len(email) > 120:
            raise ValueError("Salah satu informasi sekolah terlalu panjang.")
        if email and ("@" not in email or email.startswith("@") or email.endswith("@")):
            raise ValueError("Alamat email sekolah tidak valid.")
        logo_file = request.files.get("logo") if request.mimetype == "multipart/form-data" else None
        logo_data = logo_file.read(2 * 1024 * 1024 + 1) if logo_file and logo_file.filename else None
        remove_logo = data.get("hapus_logo") == "1"
        logo_mime = ""
        if logo_data:
            if len(logo_data) > 2 * 1024 * 1024:
                raise ValueError("Ukuran logo maksimal 2 MB.")
            if logo_data.startswith(b"\x89PNG\r\n\x1a\n"):
                logo_mime = "image/png"
            elif logo_data.startswith(b"\xff\xd8\xff"):
                logo_mime = "image/jpeg"
            elif len(logo_data) >= 12 and logo_data[:4] == b"RIFF" and logo_data[8:12] == b"WEBP":
                logo_mime = "image/webp"
            else:
                raise ValueError("Format logo harus PNG, JPEG, atau WebP.")
        if logo_data:
            get_db().execute(
                "UPDATE school_settings SET nama = ?, npsn = ?, alamat = ?, telepon = ?, email = ?, logo = ?, logo_mime = ?, updated_at = CURRENT_TIMESTAMP WHERE id = 1",
                (name, npsn, address, phone, email, logo_data, logo_mime),
            )
        elif remove_logo:
            get_db().execute(
                "UPDATE school_settings SET nama = ?, npsn = ?, alamat = ?, telepon = ?, email = ?, logo = NULL, logo_mime = '', updated_at = CURRENT_TIMESTAMP WHERE id = 1",
                (name, npsn, address, phone, email),
            )
        else:
            get_db().execute(
                "UPDATE school_settings SET nama = ?, npsn = ?, alamat = ?, telepon = ?, email = ?, updated_at = CURRENT_TIMESTAMP WHERE id = 1",
                (name, npsn, address, phone, email),
            )
        get_db().commit()
        return jsonify(settings_json())
    except ValueError as error:
        return jsonify(error=str(error)), 400


@app.get("/api/backup")
@login_required
@roles_required("admin")
def download_backup():
    snapshot = sqlite3.connect(":memory:")
    try:
        get_db().commit()
        get_db().backup(snapshot)
        contents = snapshot.serialize()
    finally:
        snapshot.close()
    filename = f"ruangsiswa-backup-{date.today().isoformat()}.sqlite3"
    return send_file(io.BytesIO(contents), mimetype="application/vnd.sqlite3", as_attachment=True, download_name=filename, max_age=0)


def validate_backup(path):
    try:
        connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        try:
            if connection.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                raise ValueError("Integritas database backup tidak valid.")
            tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
            required = {"students", "users", "grades", "attendance", "violations", "payments", "school_settings", "teacher_assignments", "classes", "teacher_profiles"}
            if not required.issubset(tables):
                raise ValueError("File tidak memiliki struktur database RuangSiswa yang lengkap.")
            required_columns = {
                "students": {"id", "nama", "nis", "kelas", "gender", "wali", "telepon", "status", "foto", "foto_mime", "created_at"},
                "users": {"id", "username", "password_hash", "role", "student_id", "active"},
                "grades": {"student_id", "mapel", "semester", "nilai", "tanggal"},
                "attendance": {"student_id", "tanggal", "waktu", "status", "catatan"},
                "violations": {"student_id", "jenis", "poin", "tanggal", "tindak_lanjut"},
                "payments": {"student_id", "tagihan", "jumlah", "tempo", "status", "dibayar"},
                "school_settings": {"id", "nama", "npsn", "alamat", "telepon", "email", "logo", "logo_mime", "auth_epoch"},
                "teacher_assignments": {"id", "teacher_id", "kelas", "mapel"},
                "classes": {"nama", "tingkat", "tahun_ajaran"},
                "teacher_profiles": {"user_id", "nama", "nip", "telepon"},
            }
            for table, expected in required_columns.items():
                actual = {row[1] for row in connection.execute(f"PRAGMA table_info({table})")}
                if not expected.issubset(actual):
                    raise ValueError("Struktur database backup tidak cocok dengan RuangSiswa.")
            if not connection.execute("SELECT 1 FROM users WHERE role = 'admin' AND active = 1 LIMIT 1").fetchone():
                raise ValueError("Database backup harus memiliki minimal satu akun admin aktif.")
            if connection.execute("PRAGMA foreign_key_check").fetchone():
                raise ValueError("Database backup memiliki relasi data yang tidak valid.")
        finally:
            connection.close()
    except sqlite3.Error as error:
        raise ValueError("File bukan database SQLite RuangSiswa yang valid.") from error


def migrate_backup(path):
    connection = sqlite3.connect(path)
    try:
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        base_tables = {"students", "users", "grades", "attendance", "violations", "payments", "school_settings"}
        if not base_tables.issubset(tables):
            raise ValueError("File tidak memiliki struktur database RuangSiswa yang lengkap.")
        student_columns = {row[1] for row in connection.execute("PRAGMA table_info(students)")}
        if "foto" not in student_columns:
            connection.execute("ALTER TABLE students ADD COLUMN foto BLOB")
        if "foto_mime" not in student_columns:
            connection.execute("ALTER TABLE students ADD COLUMN foto_mime TEXT NOT NULL DEFAULT ''")
        grade_columns = {row[1] for row in connection.execute("PRAGMA table_info(grades)")}
        if "jenis" not in grade_columns:
            connection.execute("ALTER TABLE grades ADD COLUMN jenis TEXT NOT NULL DEFAULT ''")
        user_columns = {row[1] for row in connection.execute("PRAGMA table_info(users)")}
        if "must_change_password" not in user_columns:
            connection.execute("ALTER TABLE users ADD COLUMN must_change_password INTEGER NOT NULL DEFAULT 0")
            connection.execute("UPDATE users SET must_change_password = 1 WHERE role = 'siswa'")
        settings_columns = {row[1] for row in connection.execute("PRAGMA table_info(school_settings)")}
        if "auth_epoch" not in settings_columns:
            connection.execute("ALTER TABLE school_settings ADD COLUMN auth_epoch TEXT NOT NULL DEFAULT ''")
        connection.execute("UPDATE school_settings SET auth_epoch = ? WHERE id = 1 AND auth_epoch = ''", (secrets.token_urlsafe(32),))
        connection.execute(
            "CREATE TABLE IF NOT EXISTS teacher_assignments ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, "
            "kelas TEXT NOT NULL, mapel TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, "
            "UNIQUE (teacher_id, kelas, mapel))"
        )
        connection.execute(
            "CREATE TABLE IF NOT EXISTS classes ("
            "nama TEXT PRIMARY KEY, tingkat TEXT NOT NULL DEFAULT '', "
            "tahun_ajaran TEXT NOT NULL DEFAULT '2025/2026', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)"
        )
        connection.execute(
            "CREATE TABLE IF NOT EXISTS teacher_profiles ("
            "user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE, nama TEXT NOT NULL, "
            "nip TEXT NOT NULL UNIQUE, telepon TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)"
        )
        connection.execute(
            "INSERT OR IGNORE INTO classes (nama, tingkat) "
            "SELECT DISTINCT kelas, CASE WHEN kelas GLOB '[0-9]*' THEN substr(kelas, 1, 1) ELSE '' END FROM students WHERE trim(kelas) != ''"
        )
        connection.execute(
            "INSERT OR IGNORE INTO classes (nama, tingkat) "
            "SELECT DISTINCT kelas, CASE WHEN kelas GLOB '[0-9]*' THEN substr(kelas, 1, 1) ELSE '' END FROM teacher_assignments WHERE trim(kelas) != ''"
        )
        connection.commit()
    except sqlite3.Error as error:
        raise ValueError("File bukan database SQLite RuangSiswa yang valid.") from error
    finally:
        connection.close()


@app.post("/api/restore")
@login_required
@roles_required("admin")
@csrf_required
def restore_backup():
    uploaded = request.files.get("database")
    if not uploaded or not uploaded.filename:
        return jsonify(error="Pilih file backup SQLite terlebih dahulu."), 400
    db_path = Path(app.config["DATABASE"])
    db_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(prefix="ruangsiswa-restore-", suffix=".sqlite3", dir=db_path.parent, delete=False) as temporary_file:
            temporary_path = Path(temporary_file.name)
            uploaded.save(temporary_file)
        migrate_backup(temporary_path)
        validate_backup(temporary_path)
        source = sqlite3.connect(f"file:{temporary_path}?mode=ro", uri=True)
        try:
            destination = get_db()
            destination.commit()
            destination.execute("PRAGMA foreign_keys = OFF")
            source.backup(destination)
            destination.execute("UPDATE school_settings SET auth_epoch = ?, updated_at = CURRENT_TIMESTAMP WHERE id = 1", (secrets.token_urlsafe(32),))
            destination.commit()
            destination.execute("PRAGMA foreign_keys = ON")
        finally:
            source.close()
        session.clear()
        return jsonify(ok=True, message="Database dipulihkan. Silakan masuk kembali.")
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.Error:
        get_db().rollback()
        return jsonify(error="Database tidak dapat dipulihkan dari file tersebut."), 400
    finally:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink()


def register_read_write_routes(table, endpoint, allowed_read, allowed_write, fields, order="id DESC"):
    @app.get(f"/api/{endpoint}", endpoint=f"get_{endpoint}")
    @login_required
    @roles_required(*allowed_read)
    def read_records():
        rows = scoped_rows(table, order)
        return jsonify(items=rows_as_json(rows, rename_student=True))

    @app.post(f"/api/{endpoint}", endpoint=f"post_{endpoint}")
    @login_required
    @roles_required(*allowed_write)
    @csrf_required
    def write_record():
        data = payload()
        try:
            values = {column: converter(data, key, label) for column, key, label, converter in fields}
            if not get_db().execute("SELECT 1 FROM students WHERE id = ?", (values["student_id"],)).fetchone():
                raise ValueError("Siswa tidak ditemukan.")
            if not teacher_can_access_student(values["student_id"], values.get("mapel")):
                return jsonify(error="Anda tidak ditugaskan mengajar kelas atau mata pelajaran siswa ini."), 403
            if table == "attendance" and get_db().execute(
                "SELECT 1 FROM attendance WHERE student_id = ? AND tanggal = ? LIMIT 1",
                (values["student_id"], values["tanggal"]),
            ).fetchone():
                return jsonify(error="Kehadiran siswa ini sudah dicatat pada tanggal tersebut."), 409
            columns = list(values)
            values["created_by"] = current_user()["id"]
            columns.append("created_by")
            placeholders = ", ".join(f":{column}" for column in columns)
            column_sql = ", ".join(columns)
            cursor = get_db().execute(f"INSERT INTO {table} ({column_sql}) VALUES ({placeholders})", values)
            get_db().commit()
            row = get_db().execute(f"SELECT * FROM {table} WHERE id = ?", (cursor.lastrowid,)).fetchone()
            return jsonify(item=rows_as_json([row], rename_student=True)[0]), 201
        except ValueError as error:
            return jsonify(error=str(error)), 400
        except sqlite3.IntegrityError as error:
            get_db().rollback()
            if table == "attendance" and "attendance already recorded" in str(error).lower():
                return jsonify(error="Kehadiran siswa ini sudah dicatat pada tanggal tersebut."), 409
            return jsonify(error="Data tidak valid."), 400


def student_value(data, _key, _label):
    return text(data, "siswa", "Siswa", 50)


def string_value(key, label, maximum=160):
    return lambda data, _unused, _label: text(data, key, label, maximum)


def enum_value(key, label, allowed):
    def convert(data, _unused, _label):
        value = text(data, key, label)
        if value not in allowed:
            raise ValueError(f"{label} tidak valid.")
        return value
    return convert


def integer_value(key, label, minimum, maximum=None):
    def convert(data, _unused, _label):
        try:
            value = int(data.get(key))
        except (TypeError, ValueError):
            raise ValueError(f"{label} harus berupa angka.") from None
        if value < minimum or (maximum is not None and value > maximum):
            raise ValueError(f"{label} di luar batas yang diizinkan.")
        return value
    return convert


def selected_violation_type(data):
    type_id = integer_value("violation_type_id", "Jenis pelanggaran", 1)(data, "violation_type_id", "Jenis pelanggaran")
    violation_type = get_db().execute("SELECT jenis, poin FROM violation_types WHERE id = ?", (type_id,)).fetchone()
    if not violation_type:
        raise ValueError("Pilih jenis pelanggaran yang sudah ditetapkan admin.")
    return violation_type


def violation_type_name(data, _key, _label):
    return selected_violation_type(data)["jenis"]


def violation_type_points(data, _key, _label):
    return selected_violation_type(data)["poin"]


@app.get("/api/violation-types")
@login_required
@roles_required("admin", "guru")
def list_violation_types():
    rows = get_db().execute("SELECT id, jenis, poin FROM violation_types ORDER BY jenis COLLATE NOCASE").fetchall()
    return jsonify(items=[dict(row) for row in rows])


@app.post("/api/violation-types")
@login_required
@roles_required("admin")
@csrf_required
def create_violation_type():
    data = payload()
    try:
        jenis = text(data, "jenis", "Jenis pelanggaran", 160)
        poin = integer_value("poin", "Poin", 1, 1000)(data, "poin", "Poin")
        cursor = get_db().execute("INSERT INTO violation_types (jenis, poin) VALUES (?, ?)", (jenis, poin))
        get_db().commit()
        return jsonify(item={"id": cursor.lastrowid, "jenis": jenis, "poin": poin}), 201
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Jenis pelanggaran tersebut sudah terdaftar."), 409


@app.put("/api/violation-types/<int:type_id>")
@login_required
@roles_required("admin")
@csrf_required
def update_violation_type(type_id):
    data = payload()
    try:
        jenis = text(data, "jenis", "Jenis pelanggaran", 160)
        poin = integer_value("poin", "Poin", 1, 1000)(data, "poin", "Poin")
        db = get_db()
        if not db.execute("SELECT 1 FROM violation_types WHERE id = ?", (type_id,)).fetchone():
            return jsonify(error="Jenis pelanggaran tidak ditemukan."), 404
        db.execute("UPDATE violation_types SET jenis = ?, poin = ? WHERE id = ?", (jenis, poin, type_id))
        db.commit()
        return jsonify(item={"id": type_id, "jenis": jenis, "poin": poin})
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Jenis pelanggaran tersebut sudah terdaftar."), 409


@app.delete("/api/violation-types/<int:type_id>")
@login_required
@roles_required("admin")
@csrf_required
def delete_violation_type(type_id):
    cursor = get_db().execute("DELETE FROM violation_types WHERE id = ?", (type_id,))
    get_db().commit()
    if not cursor.rowcount:
        return jsonify(error="Jenis pelanggaran tidak ditemukan."), 404
    return jsonify(ok=True)


def date_value(key, label):
    def convert(data, _unused, _label):
        value = text(data, key, label, 10)
        valid_date(value, label)
        return value
    return convert


def time_value(data, _key, _label):
    value = text(data, "waktu", "Waktu", 5)
    try:
        datetime.strptime(value, "%H:%M")
    except ValueError:
        raise ValueError("Waktu harus menggunakan format JJ:MM.") from None
    return value


@app.post("/api/grades/batch")
@login_required
@roles_required("admin", "guru")
@csrf_required
def create_grade_batch():
    data = payload()
    items = data.get("items")
    if not isinstance(items, list) or not items:
        return jsonify(error="Tambahkan minimal satu rincian nilai."), 400
    if len(items) > 20:
        return jsonify(error="Maksimal 20 nilai dalam satu pengiriman."), 400
    try:
        student_id = student_value(data, "siswa", "Siswa")
        subject = string_value("mapel", "Mata pelajaran", 120)(data, "mapel", "Mata pelajaran")
        semester = enum_value("semester", "Semester", {"Ganjil", "Genap"})(data, "semester", "Semester")
        tanggal = date_value("tanggal", "Tanggal penilaian")(data, "tanggal", "Tanggal penilaian")
        db = get_db()
        if not db.execute("SELECT 1 FROM students WHERE id = ?", (student_id,)).fetchone():
            raise ValueError("Siswa tidak ditemukan.")
        if not teacher_can_access_student(student_id, subject):
            return jsonify(error="Anda tidak ditugaskan mengajar kelas atau mata pelajaran siswa ini."), 403

        parsed_items = []
        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                raise ValueError(f"Baris nilai {index} tidak valid.")
            kind = text(item, "jenis", "Jenis penilaian", 20).upper()
            if kind not in {"NH", "NH1", "NH2", "NH-1", "NH-2", "NH-3", "NH-4", "NH-5", "US"}:
                raise ValueError(f"Jenis penilaian baris {index} harus NH-1 sampai NH-5 atau US.")
            score = integer_value("nilai", "Nilai", 0, 100)(item, "nilai", "Nilai")
            parsed_items.append((kind, score))

        db.executemany(
            "INSERT INTO grades (student_id, mapel, semester, jenis, nilai, tanggal, created_by) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            [(student_id, subject, semester, kind, score, tanggal, current_user()["id"]) for kind, score in parsed_items],
        )
        db.commit()
        return jsonify(ok=True, saved=len(parsed_items)), 201
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError:
        get_db().rollback()
        return jsonify(error="Data nilai tidak valid."), 400


@app.put("/api/grades/<int:grade_id>")
@login_required
@roles_required("admin", "guru")
@csrf_required
def update_grade(grade_id):
    data = payload()
    try:
        kind = text(data, "jenis", "Jenis penilaian", 20).upper()
        allowed_types = {"NH", "NH1", "NH2", "NH-1", "NH-2", "NH-3", "NH-4", "NH-5", "US"}
        if kind not in allowed_types:
            raise ValueError("Jenis penilaian harus NH-1 sampai NH-5 atau US.")
        score = integer_value("nilai", "Nilai", 0, 100)(data, "nilai", "Nilai")
        db = get_db()
        existing = db.execute("SELECT student_id, mapel FROM grades WHERE id = ?", (grade_id,)).fetchone()
        if not existing:
            return jsonify(error="Rincian nilai tidak ditemukan."), 404
        if not teacher_can_access_student(existing["student_id"], existing["mapel"]):
            return jsonify(error="Anda tidak ditugaskan mengajar kelas atau mata pelajaran nilai ini."), 403
        db.execute("UPDATE grades SET jenis = ?, nilai = ? WHERE id = ?", (kind, score, grade_id))
        db.commit()
        updated = db.execute("SELECT * FROM grades WHERE id = ?", (grade_id,)).fetchone()
        return jsonify(item=rows_as_json([updated], rename_student=True)[0])
    except ValueError as error:
        return jsonify(error=str(error)), 400


register_read_write_routes(
    "grades", "grades", ("admin", "guru", "siswa"), ("admin", "guru"),
    [("student_id", "siswa", "Siswa", student_value), ("mapel", "mapel", "Mata pelajaran", string_value("mapel", "Mata pelajaran")), ("semester", "semester", "Semester", enum_value("semester", "Semester", {"Ganjil", "Genap"})), ("nilai", "nilai", "Nilai", integer_value("nilai", "Nilai", 0, 100)), ("tanggal", "tanggal", "Tanggal", date_value("tanggal", "Tanggal"))],
)


@app.post("/api/attendance/batch")
@login_required
@roles_required("admin", "guru")
@csrf_required
def create_attendance_batch():
    data = payload()
    items = data.get("items")
    if not isinstance(items, list) or not items:
        return jsonify(error="Pilih minimal satu siswa untuk dicatat."), 400
    if len(items) > 500:
        return jsonify(error="Maksimal 500 siswa dalam satu pencatatan."), 400
    try:
        class_name = text(data, "kelas", "Kelas", 80)
        attendance_date = date_value("tanggal", "Tanggal")(data, "tanggal", "Tanggal")
        parsed_items = []
        seen_students = set()
        db = get_db()
        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                raise ValueError(f"Data kehadiran baris {index} tidak valid.")
            student_id = student_value(item, "siswa", "Siswa")
            if student_id in seen_students:
                raise ValueError(f"Siswa pada baris {index} tercatat lebih dari sekali.")
            seen_students.add(student_id)
            student = db.execute("SELECT kelas FROM students WHERE id = ?", (student_id,)).fetchone()
            if not student or student["kelas"] != class_name:
                raise ValueError(f"Siswa pada baris {index} bukan anggota kelas yang dipilih.")
            if db.execute(
                "SELECT 1 FROM attendance WHERE student_id = ? AND tanggal = ? LIMIT 1",
                (student_id, attendance_date),
            ).fetchone():
                return jsonify(error=f"Kehadiran {student_id} sudah dicatat pada tanggal tersebut."), 409
            if not teacher_can_access_student(student_id):
                return jsonify(error="Anda tidak ditugaskan mengajar kelas siswa ini."), 403
            status = enum_value("status", "Status kehadiran", {"Hadir", "Izin", "Sakit", "Alpa"})(item, "status", "Status kehadiran")
            note = str(item.get("catatan", "")).strip()[:300]
            parsed_items.append((student_id, attendance_date, "", status, note, current_user()["id"]))

        db.executemany(
            "INSERT INTO attendance (student_id, tanggal, waktu, status, catatan, created_by) VALUES (?, ?, ?, ?, ?, ?)",
            parsed_items,
        )
        db.commit()
        return jsonify(ok=True, saved=len(parsed_items)), 201
    except ValueError as error:
        get_db().rollback()
        return jsonify(error=str(error)), 400
    except sqlite3.IntegrityError as error:
        get_db().rollback()
        if "attendance already recorded" in str(error).lower():
            return jsonify(error="Kehadiran salah satu siswa sudah dicatat pada tanggal tersebut."), 409
        return jsonify(error="Data kehadiran tidak valid."), 400


register_read_write_routes(
    "attendance", "attendance", ("admin", "guru", "siswa"), ("admin", "guru"),
    [("student_id", "siswa", "Siswa", student_value), ("tanggal", "tanggal", "Tanggal", date_value("tanggal", "Tanggal")), ("waktu", "waktu", "Waktu", time_value), ("status", "status", "Status", enum_value("status", "Status kehadiran", {"Hadir", "Izin", "Sakit", "Alpa"})), ("catatan", "catatan", "Keterangan", lambda data, _key, _label: str(data.get("catatan", "")).strip()[:300])],
    "tanggal DESC, id DESC",
)
register_read_write_routes(
    "violations", "violations", ("admin", "guru", "siswa"), ("admin", "guru"),
    [("student_id", "siswa", "Siswa", student_value), ("jenis", "violation_type_id", "Jenis pelanggaran", violation_type_name), ("poin", "violation_type_id", "Poin", violation_type_points), ("tanggal", "tanggal", "Tanggal", date_value("tanggal", "Tanggal")), ("tindak_lanjut", "tindakLanjut", "Tindak lanjut", enum_value("tindakLanjut", "Tindak lanjut", {"Ditindaklanjuti", "Perlu tindak lanjut"}))],
    "tanggal DESC, id DESC",
)
register_read_write_routes(
    "payments", "payments", ("admin", "siswa"), ("admin",),
    [("student_id", "siswa", "Siswa", student_value), ("tagihan", "tagihan", "Jenis tagihan", string_value("tagihan", "Jenis tagihan")), ("jumlah", "jumlah", "Jumlah", integer_value("jumlah", "Jumlah", 1)), ("tempo", "tempo", "Jatuh tempo", date_value("tempo", "Jatuh tempo")), ("status", "status", "Status", enum_value("status", "Status pembayaran", {"Lunas", "Menunggu", "Terlambat"})), ("dibayar", "dibayar", "Tanggal pembayaran", lambda data, _key, _label: (str(data.get("dibayar") or "") if data.get("status") == "Lunas" else ""))],
    "tempo ASC, id DESC",
)


@app.get("/api/dashboard")
@login_required
def dashboard():
    user = current_user()
    student_filter = " WHERE student_id = ?" if user["role"] == "siswa" else ""
    params = (user["student_id"],) if student_filter else ()
    db = get_db()
    student_count = db.execute("SELECT COUNT(*) FROM students" + (" WHERE id = ?" if student_filter else ""), params).fetchone()[0]
    grade_avg = db.execute("SELECT ROUND(AVG(nilai)) FROM grades" + student_filter, params).fetchone()[0] or 0
    today_count = db.execute("SELECT COUNT(*) FROM attendance" + student_filter + (" AND" if student_filter else " WHERE") + " tanggal = ? AND status = 'Hadir'", (*params, date.today().isoformat())).fetchone()[0]
    result = {"student_count": student_count, "grade_average": grade_avg, "present_today": today_count}
    if user["role"] != "guru":
        paid_sql = "SELECT COALESCE(SUM(jumlah), 0) FROM payments WHERE status = 'Lunas'"
        due_sql = "SELECT COALESCE(SUM(jumlah), 0) FROM payments WHERE status != 'Lunas'"
        if student_filter:
            paid_sql += " AND student_id = ?"
            due_sql += " AND student_id = ?"
        result["paid_total"] = db.execute(paid_sql, params).fetchone()[0]
        result["due_total"] = db.execute(due_sql, params).fetchone()[0]
    return jsonify(result)


def seed_demo():
    db = get_db()
    today = date.today()
    def past(days):
        from datetime import timedelta
        return (today - timedelta(days=days)).isoformat()
    def future(days):
        from datetime import timedelta
        return (today + timedelta(days=days)).isoformat()

    students = [
        ("S1", "Aditya Pratama", "24001", "XII IPA 1", "Laki-laki", "Dewi Pratama", "0812-3456-7890"),
        ("S2", "Nadia Putri", "24002", "XI IPS 2", "Perempuan", "Rudi Santoso", "0813-2345-6789"),
        ("S3", "Raka Mahendra", "24003", "X IPA 1", "Laki-laki", "Sinta Mahendra", "0821-5678-9012"),
        ("S4", "Salsabila Rahman", "24004", "XII IPS 1", "Perempuan", "Ahmad Rahman", "0812-7654-3210"),
        ("S5", "Fajar Nugraha", "24005", "XI IPA 2", "Laki-laki", "Maya Nugraha", "0852-1111-2233"),
    ]
    for record in students:
        db.execute("INSERT OR IGNORE INTO students (id, nama, nis, kelas, gender, wali, telepon) VALUES (?, ?, ?, ?, ?, ?, ?)", record)
        db.execute(
            "INSERT OR IGNORE INTO classes (nama, tingkat) VALUES (?, ?)",
            (record[3], infer_class_level(record[3])),
        )
    accounts = [("admin", "admin12345", "admin", None), ("guru", "guru12345", "guru", None), ("siswa", "siswa12345", "siswa", "S1")]
    for username, password, role, student_id in accounts:
        exists = db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone()
        if not exists:
            db.execute("INSERT INTO users (username, password_hash, role, student_id, must_change_password) VALUES (?, ?, ?, ?, ?)", (username, generate_password_hash(password), role, student_id, int(role == "siswa")))
    admin_id = db.execute("SELECT id FROM users WHERE username = 'admin'").fetchone()[0]
    teacher_id = db.execute("SELECT id FROM users WHERE username = 'guru'").fetchone()[0]
    assignments = [
        ("XII IPA 1", "Kimia"), ("XII IPA 1", "Matematika"),
        ("XI IPS 2", "Bahasa Indonesia"), ("X IPA 1", "Fisika"),
        ("XII IPS 1", "Ekonomi"), ("XI IPA 2", "Biologi"),
    ]
    for class_name, subject in assignments:
        db.execute(
            "INSERT OR IGNORE INTO teacher_assignments (teacher_id, kelas, mapel) VALUES (?, ?, ?)",
            (teacher_id, class_name, subject),
        )
    grades = [("S1", "Matematika", "Ganjil", 92, past(1)), ("S2", "Bahasa Indonesia", "Ganjil", 87, past(1)), ("S3", "Fisika", "Ganjil", 78, past(2)), ("S4", "Ekonomi", "Ganjil", 90, past(2))]
    for record in grades:
        db.execute("INSERT INTO grades (student_id, mapel, semester, nilai, tanggal, created_by) SELECT ?, ?, ?, ?, ?, ? WHERE NOT EXISTS (SELECT 1 FROM grades WHERE student_id = ? AND mapel = ?)", (*record, teacher_id, record[0], record[1]))
    attendance = [("S1", today.isoformat(), "07:02", "Hadir", ""), ("S2", today.isoformat(), "07:05", "Hadir", ""), ("S3", today.isoformat(), "07:11", "Hadir", ""), ("S4", today.isoformat(), "07:08", "Izin", "Keperluan keluarga"), ("S5", today.isoformat(), "07:14", "Hadir", ""), ("S1", past(1), "07:03", "Hadir", "")]
    for record in attendance:
        db.execute("INSERT INTO attendance (student_id, tanggal, waktu, status, catatan, created_by) SELECT ?, ?, ?, ?, ?, ? WHERE NOT EXISTS (SELECT 1 FROM attendance WHERE student_id = ? AND tanggal = ?)", (*record, teacher_id, record[0], record[1]))
    violations = [("S5", "Terlambat masuk sekolah", 5, past(1), "Ditindaklanjuti"), ("S3", "Tidak memakai atribut lengkap", 3, past(3), "Perlu tindak lanjut")]
    for jenis, poin in (("Terlambat masuk sekolah", 5), ("Tidak memakai atribut lengkap", 3)):
        db.execute("INSERT OR IGNORE INTO violation_types (jenis, poin) VALUES (?, ?)", (jenis, poin))
    for record in violations:
        db.execute("INSERT INTO violations (student_id, jenis, poin, tanggal, tindak_lanjut, created_by) SELECT ?, ?, ?, ?, ?, ? WHERE NOT EXISTS (SELECT 1 FROM violations WHERE student_id = ? AND jenis = ?)", (*record, teacher_id, record[0], record[1]))
    payments = [("S1", "SPP Bulanan", 350000, future(5), "Lunas", past(2)), ("S2", "SPP Bulanan", 350000, future(5), "Lunas", past(1)), ("S3", "SPP Bulanan", 350000, future(5), "Menunggu", ""), ("S4", "Buku paket", 225000, past(2), "Terlambat", ""), ("S5", "Kegiatan sekolah", 150000, future(8), "Menunggu", "")]
    for record in payments:
        db.execute("INSERT INTO payments (student_id, tagihan, jumlah, tempo, status, dibayar, created_by) SELECT ?, ?, ?, ?, ?, ?, ? WHERE NOT EXISTS (SELECT 1 FROM payments WHERE student_id = ? AND tagihan = ?)", (*record, admin_id, record[0], record[1]))
    db.commit()


def create_initial_admin():
    username = input("Nama pengguna admin: ").strip()
    password = getpass.getpass("Kata sandi (minimal 8 karakter): ")
    confirmation = getpass.getpass("Ulangi kata sandi: ")
    if len(password) < 8 or password != confirmation:
        raise SystemExit("Kata sandi tidak cocok atau kurang dari 8 karakter.")
    try:
        get_db().execute("INSERT INTO users (username, password_hash, role) VALUES (?, ?, 'admin')", (username, generate_password_hash(password)))
        get_db().commit()
    except sqlite3.IntegrityError:
        raise SystemExit("Nama pengguna sudah digunakan.") from None
    print("Akun admin berhasil dibuat.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="RuangSiswa - sistem informasi sekolah")
    parser.add_argument("--seed-demo", action="store_true", help="Isi database dengan data dan akun demo")
    parser.add_argument("--create-admin", action="store_true", help="Buat akun admin pertama secara interaktif")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5000)
    args = parser.parse_args()
    with app.app_context():
        init_db()
        if args.seed_demo:
            seed_demo()
            print("Data contoh siap. Akun: admin, guru, siswa (kata sandi tercantum di README).")
        if args.create_admin:
            create_initial_admin()
    app.run(host=args.host, port=args.port, debug=False)